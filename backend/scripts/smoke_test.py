"""Manual smoke test for backend API endpoints."""

from __future__ import annotations

import io
import sys

from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services import fetch_service
from app.services import pipeline as pipeline_module

SAMPLE_HTML = """
<html><head><title>Test Story</title></head>
<body><article>
<p>The government announced a sweeping new national budget on Tuesday for all citizens across the country today.</p>
<p>Critics said the budget disproportionately favours wealthy households and large corporations in major cities.</p>
</article></body></html>
"""

PAYWALL_HTML = "<html><body>Subscribe to continue reading this article.</body></html>"


class FakeEmbeddingService:
    model_name = "fake-model"

    def encode_paragraph_chunks(self, chunks):
        results = []
        for chunk in chunks:
            text = str(chunk.get("text", "")).strip()
            if not text:
                continue
            results.append(
                {
                    **chunk,
                    "embedding": [0.1, 0.2, 0.3],
                    "dimension": 3,
                    "model_name": self.model_name,
                }
            )
        return results


async def fake_fetch_html(url, settings, *, article_ref=None, progress=None):
    return SAMPLE_HTML


async def paywall_fetch_html(url, settings, *, article_ref=None, progress=None):
    return PAYWALL_HTML


def main() -> int:
    pipeline_module.get_embedding_service = lambda: FakeEmbeddingService()
    fetch_service.fetch_html = fake_fetch_html
    client = TestClient(app)
    results: list[tuple[str, bool, str]] = []

    def check(name: str, ok: bool, detail: str = "") -> None:
        results.append((name, ok, detail))
        status = "PASS" if ok else "FAIL"
        suffix = f" - {detail}" if detail else ""
        print(f"[{status}] {name}{suffix}")

    r = client.get("/health")
    check("GET /health", r.status_code == 200 and r.json().get("status") == "ok", str(r.json()))

    r = client.get("/health/db")
    check("GET /health/db", r.status_code == 200 and "database" in r.json(), str(r.json()))

    r = client.post("/api/fetch", json={"url": "https://example.com/story"})
    data = r.json()
    check(
        "POST /api/fetch",
        r.status_code == 200 and bool(data.get("body_text")),
        f"elapsed={data.get('processing', {}).get('total_elapsed_seconds')}",
    )

    r = client.post("/api/fetch", json={"url": "ftp://bad.com"})
    check(
        "POST /api/fetch invalid URL",
        r.status_code == 422 and r.json()["detail"].get("code") == "url_invalid_scheme",
    )

    doc = Document()
    doc.add_paragraph(
        "The government announced a sweeping new national budget on Tuesday for all citizens."
    )
    buf = io.BytesIO()
    doc.save(buf)
    r = client.post(
        "/api/upload",
        files={
            "file": (
                "test.docx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    check(
        "POST /api/upload docx",
        r.status_code == 200 and r.json()["article"]["source_type"] == "upload",
    )

    r = client.post(
        "/api/compare",
        json={"article_a_url": "https://a.com/x", "article_b_url": "https://b.com/y"},
    )
    compare_data = r.json()
    check(
        "POST /api/compare",
        r.status_code == 200 and len(compare_data["articles"]) == 2 and compare_data.get("processing"),
        f"errors={len(compare_data['errors'])}",
    )

    doc2 = Document()
    doc2.add_paragraph(
        "Opposition leaders argued that essential services for low-income families had been neglected again today."
    )
    buf2 = io.BytesIO()
    doc2.save(buf2)
    r = client.post(
        "/api/compare/files",
        data={"article_b_url": "https://b.com/y"},
        files={
            "article_a_file": (
                "a.docx",
                buf2.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    check("POST /api/compare/files", r.status_code == 200 and len(r.json()["articles"]) == 2)

    with client.stream(
        "POST",
        "/api/compare/stream",
        json={"article_a_url": "https://a.com/x", "article_b_url": "https://b.com/y"},
    ) as stream_response:
        body = stream_response.read().decode("utf-8")
    progress_events = body.count("event: progress")
    has_result = "event: result" in body
    check(
        "POST /api/compare/stream SSE",
        stream_response.status_code == 200 and progress_events > 0 and has_result,
        f"progress_events={progress_events}",
    )

    fetch_service.fetch_html = paywall_fetch_html
    r = client.post("/api/fetch", json={"url": "https://paywall.example.com/story"})
    check(
        "Paywall error code",
        r.status_code == 422 and r.json()["detail"].get("code") == "extraction_paywall",
    )

    failed = [item for item in results if not item[1]]
    print("---")
    print(f"Summary: {len(results) - len(failed)}/{len(results)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
