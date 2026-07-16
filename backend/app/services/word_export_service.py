"""Build downloadable Word (.docx) documents from cleaned text."""

from __future__ import annotations

import io
import re

from docx import Document


def _safe_filename(name: str, *, default: str = "converted", suffix: str = ".docx") -> str:
    """Turn an arbitrary title/filename into a safe .docx download name."""
    stem = re.sub(r"[^\w\-. ]+", "_", (name or default).strip()) or default
    stem = stem.rsplit(".", 1)[0] if stem.lower().endswith((".pdf", ".docx")) else stem
    stem = stem.strip(" ._") or default
    return f"{stem}{suffix}"


def build_docx(body_text: str, *, title: str | None = None) -> bytes:
    """Create a .docx file (as bytes) from title + paragraph text."""
    document = Document()

    if title:
        document.add_heading(title, level=1)

    for block in re.split(r"\n\s*\n", body_text.strip()):
        block = block.strip()
        if not block:
            continue
        for line in block.splitlines():
            line = line.strip()
            if line:
                document.add_paragraph(line)

    if title:
        document.core_properties.title = title

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def download_filename_for(source_filename: str | None, title: str | None) -> str:
    """Pick a friendly .docx download name from the source file or title."""
    base = title or source_filename or "converted"
    return _safe_filename(base)
