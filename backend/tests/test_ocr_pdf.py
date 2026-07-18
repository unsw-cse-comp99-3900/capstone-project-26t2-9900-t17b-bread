"""Tests for PDF type detection, OCR cleaning, Word export, and OCR endpoints."""

from __future__ import annotations

import io

import fitz  # PyMuPDF
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services import ocr_service
from app.services.errors import ErrorCode
from app.services.ocr_cleaning import clean_ocr_pages
from app.services.pdf_analysis import analyze_pdf
from app.services.word_export_service import build_docx, download_filename_for

client = TestClient(app)

_DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

BODY = (
    "The government announced a sweeping new national budget on Tuesday for all citizens. "
    "Critics said the plan favours wealthy households and large corporations across the country."
)


def _make_text_pdf(text: str = BODY) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=12)
    buffer = io.BytesIO()
    doc.save(buffer)
    doc.close()
    return buffer.getvalue()


def _make_image_pdf() -> bytes:
    """A PDF whose page is an image with no selectable text (a fake 'scan')."""
    doc = fitz.open()
    page = doc.new_page(width=300, height=200)
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 100))
    pix.set_rect(pix.irect, (255, 255, 255))
    page.insert_image(fitz.Rect(20, 20, 220, 120), pixmap=pix)
    buffer = io.BytesIO()
    doc.save(buffer)
    doc.close()
    return buffer.getvalue()


# --- PDF type detection -----------------------------------------------------

def test_detect_text_pdf():
    analysis = analyze_pdf(_make_text_pdf())
    assert analysis.pdf_type == "text"
    assert not analysis.is_image_based
    assert "budget" in analysis.extracted_text.lower()


def test_detect_image_pdf():
    analysis = analyze_pdf(_make_image_pdf())
    assert analysis.pdf_type == "image"
    assert analysis.is_image_based
    assert analysis.pages_with_images >= 1


def test_pdf_type_endpoint_text():
    r = client.post(
        "/api/upload/pdf-type",
        files={"file": ("doc.pdf", _make_text_pdf(), "application/pdf")},
    )
    assert r.status_code == 200
    data = r.json()
    assert data["pdf_type"] == "text"
    assert data["is_image_based"] is False
    assert data["recommended_endpoint"] == "/api/upload"


def test_pdf_type_endpoint_image():
    r = client.post(
        "/api/upload/pdf-type",
        files={"file": ("scan.pdf", _make_image_pdf(), "application/pdf")},
    )
    assert r.status_code == 200
    data = r.json()
    assert data["pdf_type"] == "image"
    assert data["recommended_endpoint"] == "/api/upload/pdf-to-word"
    assert "ocr_available" in data


# --- OCR noise cleaning -----------------------------------------------------

def test_clean_ocr_removes_page_numbers_and_headers():
    pages = [
        "DAILY NEWS\nThe budget passed today after a long debate in parliament.\n1",
        "DAILY NEWS\nOfficials defended the plan as fair and balanced for everyone.\n2",
        "DAILY NEWS\nCritics argued it favours the wealthy over ordinary families.\n3",
    ]
    cleaned = clean_ocr_pages(pages)
    assert "DAILY NEWS" not in cleaned  # repeated header removed
    assert "budget passed today" in cleaned
    assert "\n1" not in cleaned and "\n2" not in cleaned  # page numbers removed


def test_clean_ocr_drops_symbol_noise():
    pages = ["|||\nReal sentence with meaningful words here.\n~ * ^"]
    cleaned = clean_ocr_pages(pages)
    assert "Real sentence with meaningful words here." in cleaned
    assert "|||" not in cleaned


def test_clean_ocr_rejoins_hyphenation():
    pages = ["The govern-\nment announced a new plan for the economy today."]
    cleaned = clean_ocr_pages(pages)
    assert "government announced" in cleaned


# --- Word export ------------------------------------------------------------

def test_build_docx_roundtrip():
    docx_bytes = build_docx("Para one.\n\nPara two.", title="My Title")
    document = Document(io.BytesIO(docx_bytes))
    texts = [p.text for p in document.paragraphs]
    assert "My Title" in texts
    assert "Para one." in texts
    assert "Para two." in texts


def test_download_filename():
    assert download_filename_for("scan.pdf", None).endswith(".docx")
    assert download_filename_for(None, "Budget Report").endswith(".docx")


# --- pdf-to-word endpoint ---------------------------------------------------

def test_pdf_to_word_text_pdf_downloads_docx():
    r = client.post(
        "/api/upload/pdf-to-word",
        files={"file": ("doc.pdf", _make_text_pdf(), "application/pdf")},
    )
    assert r.status_code == 200
    assert r.headers["content-type"] == _DOCX_MEDIA_TYPE
    assert "attachment" in r.headers.get("content-disposition", "")
    # Returned bytes should be a valid .docx.
    document = Document(io.BytesIO(r.content))
    joined = " ".join(p.text for p in document.paragraphs)
    assert "budget" in joined.lower()


def test_pdf_to_word_image_pdf_without_tesseract_reports_ocr_error(monkeypatch):
    monkeypatch.setattr(ocr_service, "_resolve_tesseract", lambda settings: None)
    r = client.post(
        "/api/upload/pdf-to-word",
        files={"file": ("scan.pdf", _make_image_pdf(), "application/pdf")},
    )
    assert r.status_code == 422
    assert r.json()["detail"]["code"] == ErrorCode.OCR_UNAVAILABLE.value


def test_pdf_to_word_rejects_non_pdf():
    r = client.post(
        "/api/upload/pdf-to-word",
        files={"file": ("notes.txt", b"hello", "text/plain")},
    )
    assert r.status_code == 422
    assert r.json()["detail"]["code"] == ErrorCode.UPLOAD_UNSUPPORTED_TYPE.value
